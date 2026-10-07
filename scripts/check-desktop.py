"""Verify the portable runtime with a real, isolated DOCX-to-PDF conversion."""
import io
import os
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ['USECTA_LOCAL_MODE'] = '1'
from docx import Document
from documents import convert_pdf, pdf_available, render_document


def main():
    if not pdf_available():
        raise RuntimeError('Bundled LibreOffice was not found in runtime/libreoffice/program.')
    sample = Document()
    sample.add_paragraph('USECTA DESKTOP SETUP CHECK')
    sample.add_paragraph('Supplier: {{ supplier_name }}')
    content = io.BytesIO()
    sample.save(content)
    rendered = render_document(content.getvalue(), {'supplier_name': 'SAMPLE ONLY'}, ['supplier_name'])
    assert Document(io.BytesIO(rendered)).paragraphs[1].text == 'Supplier: SAMPLE ONLY'
    pdf = convert_pdf(rendered)
    if not pdf.startswith(b'%PDF-'):
        raise RuntimeError('LibreOffice did not produce a valid PDF.')
    print('Portable Python imports, DOCX rendering, and LibreOffice PDF generation passed.')


if __name__ == '__main__':
    main()
