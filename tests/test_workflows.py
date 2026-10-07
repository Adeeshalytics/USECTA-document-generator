import io
import json
import sqlite3
import zipfile
from datetime import date
from pathlib import Path

import pytest
from docx import Document
from lxml import etree
from openpyxl import Workbook
from streamlit.testing.v1 import AppTest

import documents
from documents import Store, render_document, template_variables
from workflows import (choose_items, create_quotation_template, generate_files, import_catalog,
                       quotation_context, quotation_jobs, tender_context, tender_jobs, validate_catalog)

APP = Path(__file__).resolve().parents[1] / "app.py"
CATALOG = APP.parent / "catalogues/SPMC_03_2026.csv"
DETAILS = {"closing_date": "2026-11-20", "closing_time": "14:30", "bid_valid_until": "2027-02-20", "tender_title": "Example tender"}
SUPPLIER = {"supplier_name": "Example Chemicals Ltd", "supplier_country": "China", "supplier_address": "100 Example Road", "supplier_registration_number": "TEST-123"}


def word(text):
    doc = Document()
    doc.add_paragraph(text)
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def context(rows):
    return tender_context("SPMC/03/2026", DETAILS, rows, SUPPLIER, "2026-11-01", "2026-11-02")


@pytest.fixture
def catalog_store(tmp_path):
    store = Store(tmp_path)
    store.save_tender("SPMC/03/2026", DETAILS)
    store.save_catalog("SPMC/03/2026", import_catalog(CATALOG.read_bytes(), ".csv"))
    return store


def test_supplied_catalogue_all_50_items_and_distinct_amounts():
    rows = import_catalog(CATALOG.read_bytes(), ".csv")
    assert [r["item_number"] for r in rows] == list(map(str, range(1, 51)))
    assert rows[1] == {"item_number": "2", "description": "Benzhexol Hydrochloride BP (Trihexyphenidyl Hydrochloride BP)", "document_fee": "500.00", "quantity_kg": "50"}
    assert rows[14]["description"] == 'Mebendazole USP (Polymorph Type "A")'
    assert rows[11]["quantity_kg"] == "16000" and rows[11]["document_fee"] == "41000.00"
    assert rows[36]["quantity_kg"] == "60000" and rows[36]["document_fee"] == "11000.00"
    assert rows[-1]["description"] == "Zinc Stearate BP"


def test_lookup_is_scoped_to_tender_and_missing_items_rejected(catalog_store):
    store = catalog_store
    store.save_tender("SPMC/02/2026", DETAILS)
    store.save_catalog("SPMC/02/2026", [{"item_number": "15", "description": "Mefenamic Acid BP", "document_fee": "1000"}])
    assert choose_items(store, "SPMC/02/2026", ["15"])[0]["description"] == "Mefenamic Acid BP"
    assert choose_items(store, "SPMC/03/2026", ["15"])[0]["description"].startswith("Mebendazole")
    with pytest.raises(ValueError, match="not in tender"):
        choose_items(store, "SPMC/03/2026", ["51"])
    with pytest.raises(ValueError, match="twice"):
        choose_items(store, "SPMC/03/2026", ["15", "15"])


@pytest.mark.parametrize("rows", [
    [{"item_number": "1", "description": "Item", "document_fee": "-1"}],
    [{"item_number": "1", "description": "Item", "document_fee": "NaN"}],
    [{"item_number": "1", "description": "Item", "document_fee": ""}],
    [{"item_number": "1", "description": "Item", "document_fee": "500"}] * 2,
])
def test_bad_catalogues_leave_existing_data_intact(catalog_store, rows):
    before = catalog_store.catalog("SPMC/03/2026")
    with pytest.raises(ValueError):
        catalog_store.save_catalog("SPMC/03/2026", rows)
    assert catalog_store.catalog("SPMC/03/2026") == before


def test_excel_import_and_date_formats():
    book = Workbook()
    sheet = book.active
    sheet.append(["item_number", "description", "document_fee", "quantity_kg"])
    sheet.append([2, "Example item", "1,000.00", 50])
    output = io.BytesIO()
    book.save(output)
    rows = import_catalog(output.getvalue(), ".xlsx")
    assert rows[0]["item_number"] == "2"
    assert rows[0]["document_fee"] == "1000.00"
    data = context(rows)
    assert data["closing_date_long"] == "20 November 2026"
    assert data["closing_date_short"] == "20.11.2026"
    assert data["closing_time"] == "2.30 p.m."
    assert data["bid_valid_until"] == "20.02.2027"


def test_per_item_and_batch_generation_and_saved_snapshot(catalog_store):
    store = catalog_store
    rows = choose_items(store, "SPMC/03/2026", ["15", "34"])
    per_item = store.add_template("Cover", word("{{ bid_reference }} {{ item_number }} {{ item_description }} {{ supplier_name }}"), ["supplier_name"], workflow="tender", scope="item")
    batch = store.add_template("Agreement", word("{{ supplier_name }} — {{ agreement_item_names }}"), ["agreement_item_names"], workflow="tender", scope="supplier")
    jobs = tender_jobs(store, [per_item, batch], context(rows))
    files = generate_files(store, jobs)
    assert len(files) == 3
    texts = ["\n".join(p.text for p in Document(io.BytesIO(raw)).paragraphs) for raw in files.values()]
    assert "15" in texts[0] and "Mebendazole" in texts[0] and "Hypromellose" not in texts[0]
    assert "34" in texts[1] and "Hypromellose" in texts[1]
    assert "Mebendazole" in texts[2] and "Hypromellose" in texts[2]
    store.save_record("Saved tender run", {"workflow": "tender", "jobs": jobs})
    store.save_catalog("SPMC/03/2026", [{"item_number": "15", "description": "Changed later", "document_fee": "500"}])
    saved = json.loads(store.records()[0]["payload"])
    assert generate_files(store, saved["jobs"]) == files


def test_single_item_fields_require_correct_scope(catalog_store):
    key = catalog_store.add_template("Incorrect scope", word("{{ item_number }}"), workflow="tender")
    with pytest.raises(ValueError, match="single-item"):
        tender_jobs(catalog_store, [key], context(choose_items(catalog_store, "SPMC/03/2026", ["1"])))


def test_quotation_is_independent_and_preserves_units(catalog_store):
    create_quotation_template(catalog_store)
    fields = {"quotation_number": "QT-00282", "customer_name": "Example Customer", "customer_address": "Example address", "currency": "USD", "document_date": "1 November 2026", "quotation_expiry_date": "8 November 2026", "bid_reference": "SPMC/03/2026"}
    data = quotation_context(fields, [{"description": "Independent product", "quantity": "17.60", "unit": "MT", "unit_price": "1550"}])
    assert "bid_reference" not in data and "tender_items" not in data
    assert data["grand_total"] == "27,280.00"
    assert data["quotation_items"][0]["number"] == 1 and data["quotation_items"][0]["unit"] == "MT"
    bad = catalog_store.add_template("Bad quotation", word("{{ bid_reference }}"), workflow="quotation")
    with pytest.raises(ValueError, match="cannot use tender"):
        quotation_jobs(catalog_store, [bad], data)


def test_malformed_placeholder_is_rejected():
    with pytest.raises(ValueError, match="opening brace"):
        template_variables(word("{ document_date }}"))


def test_untouched_parts_and_section_settings_are_preserved():
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "FIXED USECTA HEADER"
    doc.sections[0].footer.paragraphs[0].text = "FIXED FOOTER"
    doc.sections[0].left_margin = 914400
    doc.add_paragraph("{{ supplier_name }}")
    stream = io.BytesIO()
    doc.save(stream)
    before = stream.getvalue()
    after = render_document(before, {"supplier_name": "Supplier & Co"})
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with zipfile.ZipFile(io.BytesIO(before)) as src, zipfile.ZipFile(io.BytesIO(after)) as dst:
        assert src.namelist() == dst.namelist()
        for name in src.namelist():
            if name != "word/document.xml":
                assert src.read(name) == dst.read(name), name
        a = etree.fromstring(src.read("word/document.xml")).find(".//w:sectPr", ns)
        b = etree.fromstring(dst.read("word/document.xml")).find(".//w:sectPr", ns)
        assert etree.tostring(a, method="c14n") == etree.tostring(b, method="c14n")


def test_placeholder_in_header_is_rendered_without_changing_styles():
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Reference {{ bid_reference }}"
    doc.add_paragraph("Fixed text")
    stream = io.BytesIO()
    doc.save(stream)
    output = render_document(stream.getvalue(), {"bid_reference": "SPMC/03/2026"})
    assert Document(io.BytesIO(output)).sections[0].header.paragraphs[0].text == "Reference SPMC/03/2026"
    with zipfile.ZipFile(io.BytesIO(stream.getvalue())) as a, zipfile.ZipFile(io.BytesIO(output)) as b:
        assert a.read("word/styles.xml") == b.read("word/styles.xml")


def test_old_database_migration_preserves_records(tmp_path):
    with sqlite3.connect(tmp_path / "records.sqlite3") as con:
        con.executescript("CREATE TABLE templates (id TEXT PRIMARY KEY, name TEXT, filename TEXT, required TEXT, created_at TEXT); CREATE TABLE records (id TEXT PRIMARY KEY, name TEXT, payload TEXT, created_at TEXT);")
        con.execute("INSERT INTO records VALUES ('old', 'Old project', '{}', '2026-01-01')")
        con.execute("INSERT INTO templates VALUES ('old', 'Original', 'old.docx', '[]', '2026-01-01')")
    store = Store(tmp_path)
    assert store.records()[0]["name"] == "Old project"
    assert store.templates()[0]["workflow"] == "legacy"
    assert store.templates()[0]["scope"] == "batch"


def test_preloaded_catalogue_preserves_user_edits(tmp_path):
    from workflows import prepare_workspace
    store = Store(tmp_path)
    prepare_workspace(store)
    assert len(store.catalog("SPMC/03/2026")) == 50
    assert store.tenders()["SPMC/03/2026"]["closing_date"] == "2026-11-03"
    assert store.tenders()["SPMC/03/2026"]["closing_time"] == "10:00"
    assert not store.tenders()["SPMC/03/2026"]["bid_valid_until"]
    store.save_catalog("SPMC/03/2026", [{"item_number": "15", "description": "Edited by user", "document_fee": "100"}])
    store.save_tender("SPMC/03/2026", DETAILS)
    prepare_workspace(store)
    assert store.catalog("SPMC/03/2026")[0]["description"] == "Edited by user"
    assert store.tenders()["SPMC/03/2026"]["closing_date"] == DETAILS["closing_date"]


def test_app_tender_lookup_generation_and_no_stale_downloads(catalog_store, monkeypatch):
    store = catalog_store
    key = store.add_template("Cover", word("{{ item_number }} {{ item_description }} {{ supplier_name }}"), ["supplier_name"], workflow="tender", scope="item")
    supplier = store.save_supplier(SUPPLIER)
    monkeypatch.setattr(documents, "Store", lambda: store)
    app = AppTest.from_file(APP, default_timeout=20).run()
    app.selectbox(key="tender_reference").set_value("SPMC/03/2026").run()
    app.multiselect(key="tender_templates").set_value([key]).run()
    app.multiselect(key="selected_items_SPMC/03/2026").set_value(["15", "34"]).run()
    app.selectbox(key="tender_supplier").set_value(supplier).run()
    app.button(key="generate_tender").click().run()
    assert not app.exception
    assert not app.error
    assert len(app.session_state["output_tender"]["files"]) == 2
    assert any("Documents generated" in x.value for x in app.success)
    app.multiselect(key="selected_items_SPMC/03/2026").set_value(["15"]).run()
    assert not any("Documents generated" in x.value for x in app.success)
    assert "field_company_name" not in [x.key for x in app.text_input]


def test_app_creates_tender_and_supplier(tmp_path, monkeypatch):
    store = Store(tmp_path)
    monkeypatch.setattr(documents, "Store", lambda: store)
    app = AppTest.from_file(APP, default_timeout=20).run()
    app.text_input(key="catalog_ref_new").set_value("SPMC/03/2026")
    app.date_input(key="catalog_date_new").set_value(date(2026, 11, 20))
    next(x for x in app.button if x.label == "Save tender").click().run()
    assert not app.exception
    assert store.tenders()["SPMC/03/2026"]["closing_date"] == "2026-11-20"
    app.text_input(key="supplier_new_supplier_name").set_value("Saved Supplier")
    next(x for x in app.button if x.label == "Save supplier").click().run()
    assert not app.exception
    assert any(s["supplier_name"] == "Saved Supplier" for s in store.suppliers().values())


def test_app_quotation_generation(tmp_path, monkeypatch):
    store = Store(tmp_path)
    monkeypatch.setattr(documents, "Store", lambda: store)
    # AppTest does not serialize data-editor edits across button reruns. Supply
    # the user-entered table at that UI boundary; rendering remains real.
    import streamlit as st
    original_editor = st.data_editor
    entered_rows = [{"description": "Independent product", "quantity": 2.0, "unit": "kg", "unit_price": 100.0}]
    monkeypatch.setattr(st, "data_editor", lambda data, **kwargs: entered_rows if kwargs.get("key") == "quotation_rows" else original_editor(data, **kwargs))
    app = AppTest.from_file(APP, default_timeout=20).run()
    app.text_input(key="quotation_number").set_value("QT-TEST")
    app.text_input(key="quote_customer_name").set_value("Example Customer")
    app.run()
    app.button(key="generate_quotation").click().run()
    assert not app.exception
    assert not app.error
    assert len(app.session_state["output_quotation"]["files"]) == 1
    output = next(iter(app.session_state["output_quotation"]["files"].values()))
    doc = Document(io.BytesIO(output))
    assert doc.tables[0].rows[1].cells[1].text == "Independent product"
    assert doc.tables[0].rows[1].cells[3].text == "kg"
    assert "SPMC" not in "\n".join(p.text for p in doc.paragraphs)
