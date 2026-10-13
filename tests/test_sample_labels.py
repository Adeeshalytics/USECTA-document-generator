import io
import json
from pathlib import Path

import pytest
from docx import Document
from streamlit.testing.v1 import AppTest

import documents
from documents import Store
from sample_labels import LABEL_GAP_POINTS, labels_document, label_files, validate_labels
from docx.oxml.ns import qn
from workflows import prepare_workspace

ROOT = Path(__file__).resolve().parents[1]


def test_two_editable_copies_per_product_and_letter_page():
    names = ["Mefanamic Acid BP 2024 / BP 2025", "Hypromellose 2910 USP 47 / USP 48",
             "Benzhexol Hydrochloride BP 2024 / BP 2025 (Trihexyphenidyl Hydrochloride BP 2024 / BP 2025)"]
    doc = Document(io.BytesIO(labels_document("SPMC/02/2026", names)))
    assert len(doc.tables) == 3
    assert doc.sections[0].page_width.pt == 612 and doc.sections[0].page_height.pt == 792
    for table, name in zip(doc.tables, names):
        row = table.rows[0]
        assert row.cells[0].text == row.cells[2].text
        assert name in row.cells[0].text
        assert "IFB NO : SPMC/02/2026" in row.cells[0].text
        assert row.cells[1].text == ""
        assert row.height is None, "Height must follow text, not a fixed label size"


def test_many_labels_flow_as_separate_boxes_with_small_gaps():
    names = [f"Product {number:03d}" for number in range(100)]
    doc = Document(io.BytesIO(labels_document("TEST-IFB", names)))
    assert len(doc.tables) == 100
    assert doc.sections[0].top_margin.pt == doc.sections[0].bottom_margin.pt == 24
    assert all("w:pageBreakBefore" not in table._tbl.xml and "w:cantSplit" in table.rows[0]._tr.xml for table in doc.tables)
    assert doc.styles["Normal"].font.size.pt == 14
    for table, name in zip(doc.tables, names):
        assert len(table.rows) == 1
        row = table.rows[0]
        assert row.height is None
        assert row.cells[0].text.count(name) == row.cells[2].text.count(name) == 1
        for cell in (row.cells[0], row.cells[2]):
            borders = cell._tc.get_or_add_tcPr().find(qn("w:tcBorders"))
            assert all(borders.find(qn("w:" + edge)).get(qn("w:val")) == "single" for edge in ("top", "bottom", "left", "right"))
    assert len(doc.paragraphs) == 100
    for gap in doc.paragraphs[:-1]:
        assert not gap.text and gap.paragraph_format.line_spacing.pt == LABEL_GAP_POINTS
        assert gap.paragraph_format.keep_with_next
        assert "w:pBdr" not in gap._p.xml


def test_previous_joined_layout_stays_available_for_saved_runs():
    names = ["First product", "Second product", "Third product"]
    original = labels_document("TEST-IFB", names, layout_version=2)
    doc = Document(io.BytesIO(original))
    assert len(doc.tables[0].rows) == 3
    assert all(row.height is None for row in doc.tables[0].rows)
    assert original == labels_document("TEST-IFB", names, layout_version=2)
    assert original != labels_document("TEST-IFB", names)


def test_saved_original_layout_remains_supported_and_deterministic():
    names = ["First product", "Second product", "Third product", "Fourth product"]
    old = label_files("TEST-IFB", names, layout_version=1)
    assert old == label_files("TEST-IFB", names, layout_version=1)
    doc = Document(io.BytesIO(next(iter(old.values()))))
    assert len(doc.tables) == 4 and doc.sections[0].top_margin.pt == 60
    assert old != label_files("TEST-IFB", names)


@pytest.mark.parametrize("ifb,names", [("", ["Name"]), ("Valid", []), ("Valid", [""]),
    ("bad\nIFB", ["Name"]), ("Valid", ["x" * 241]), ("Valid", ["bad\x00name"]), ("Valid", ["Name"] * 101)])
def test_invalid_label_input_rejected(ifb, names):
    with pytest.raises(ValueError):
        validate_labels(ifb, names)


def test_labels_escape_xml_and_reject_unknown_saved_layout():
    files = label_files("SPMC/03/2026", ["Sample <A> & B"])
    assert len(files) == 1 and "/" not in next(iter(files))
    doc = Document(io.BytesIO(next(iter(files.values()))))
    assert "Sample <A> & B" in doc.tables[0].cell(0, 0).text
    with pytest.raises(ValueError, match="not supported"):
        label_files("IFB", ["Name"], layout_version=999)


@pytest.fixture
def workspace(tmp_path, monkeypatch):
    store = Store(tmp_path)
    prepare_workspace(store)
    monkeypatch.setattr(documents, "Store", lambda: store)
    monkeypatch.setattr(documents, "pdf_available", lambda: None)
    return store


def application():
    return AppTest.from_file(ROOT / "app.py", default_timeout=20).run()


def test_manual_generation_snapshot_regeneration_and_stale_downloads(workspace):
    app = application()
    app.radio(key="page").set_value("Sample labels").run()
    assert app.button(key="labels_generate").disabled
    app.text_input(key="labels_ifb").set_value("IFB/2026-NEW").run()
    app.text_area(key="labels_names").set_value("First product\nSecond product").run()
    app.button(key="labels_generate").click().run()
    assert not app.exception and not app.error
    original = app.session_state["labels_output"]["files"]
    payload = json.loads(workspace.records()[0]["payload"])
    assert payload["product_names"] == ["First product", "Second product"]
    assert list((workspace.root / "exports").rglob("*.docx"))
    app.text_area(key="labels_names").set_value("Changed product").run()
    assert not any("Sample labels generated" in element.value for element in app.success)
    app.selectbox(key="history").set_value(workspace.records()[0]["id"]).run()
    next(button for button in app.button if button.label == "Regenerate saved run").click().run()
    assert not app.exception
    assert app.session_state["history_output"]["files"] == original


def test_catalogue_load_keeps_other_catalogue_and_draft_navigation(workspace):
    before = workspace.catalog("SPMC/03/2026")
    app = application()
    app.radio(key="page").set_value("Sample labels").run()
    app.selectbox(key="labels_reference").set_value("SPMC/03/2026").run()
    app.multiselect(key="labels_items_SPMC/03/2026").set_value(["2", "15"]).run()
    app.button(key="labels_load").click().run()
    assert not app.exception
    assert app.text_input(key="labels_ifb").value == "SPMC/03/2026"
    assert "Benzhexol" in app.text_area(key="labels_names").value
    app.text_area(key="labels_names").set_value("Exact label name BP 2025").run()
    app.radio(key="page").set_value("Suppliers").run()
    app.radio(key="page").set_value("Sample labels").run()
    assert not app.exception
    assert app.text_area(key="labels_names").value == "Exact label name BP 2025"
    assert workspace.catalog("SPMC/03/2026") == before


def test_label_pdf_failure_retains_word(workspace, monkeypatch):
    monkeypatch.setattr(documents, "pdf_available", lambda: "converter")
    monkeypatch.setattr(documents, "convert_pdfs", lambda *_: (_ for _ in ()).throw(ValueError("Test failure")))
    app = application()
    app.radio(key="page").set_value("Sample labels").run()
    app.text_input(key="labels_ifb").set_value("TEST-IFB").run()
    app.text_area(key="labels_names").set_value("Sample product").run()
    app.button(key="labels_generate").click().run()
    assert not app.exception and not app.error
    assert all(name.endswith(".docx") for name in app.session_state["labels_output"]["files"])
    assert any("Word labels were saved" in warning.value for warning in app.warning)
