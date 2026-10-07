import io
import json
from pathlib import Path
from urllib.parse import parse_qs, urlsplit

import pytest
from docx import Document
from streamlit.testing.v1 import AppTest

from cloud_store import SupabaseAPI, SupabaseError, SupabaseStore
from documents import render_document


class FakeServer:
    """Contract double for PostgREST and private storage, independent of local cache."""
    def __init__(self):
        self.rows, self.files = {}, {}

    def request(self, method, path, payload=None, raw=False, headers=None):
        if path.startswith('/rest/v1/'):
            query = parse_qs(urlsplit(path).query)
            if method == 'POST':
                key = (payload['kind'], payload['key'])
                if key in self.rows and 'merge-duplicates' not in (headers or {}).get('Prefer', ''):
                    raise SupabaseError(409)
                self.rows[key] = json.loads(json.dumps(payload['data']))
                return None
            kind = query['kind'][0][3:]
            key = query.get('key', [None])[0]
            result = [{'key': k, 'data': json.loads(json.dumps(value))} for (group, k), value in self.rows.items()
                      if group == kind and (key is None or k == key[3:])]
            offset = int(query['offset'][0])
            return sorted(result, key=lambda r: r['key'])[offset:offset + 1000]
        name = path.rsplit('/', 1)[-1]
        if method == 'POST':
            if name in self.files:
                raise SupabaseError(409)
            self.files[name] = payload
            return None
        return self.files[name]


def test_cloud_survives_empty_cache_and_preserves_template(tmp_path):
    server = FakeServer()
    store = SupabaseStore(server, cache_root=tmp_path / 'first')
    store.save_tender('SPMC/03/2026', {'closing_date': '2026-11-03', 'closing_time': '10:00', 'bid_valid_until': ''})
    rows = [{'item_number': '32', 'description': 'Di-Basic Calcium Phosphate', 'document_fee': '2000', 'quantity_kg': '4000'}]
    store.save_catalog('SPMC/03/2026', rows)
    with pytest.raises(ValueError):
        store.save_catalog('SPMC/03/2026', rows + [rows[0]])
    assert store.catalog('SPMC/03/2026')[0]['description'] == rows[0]['description']
    supplier = store.save_supplier({'supplier_name': 'Example Supplier'})
    doc = Document()
    run = doc.add_paragraph().add_run('{{ supplier_name | upper }}')
    run.bold = True
    out = io.BytesIO()
    doc.save(out)
    key = store.add_template('Agreement', out.getvalue(), ['supplier_name'], workflow='tender', scope='supplier')
    record = store.save_record('Original', {'supplier': 'Example Supplier'})
    store.save_supplier({'supplier_name': 'Changed'}, supplier)
    restarted = SupabaseStore(server, cache_root=tmp_path / 'empty-cache')
    assert restarted.suppliers()[supplier]['supplier_name'] == 'Changed'
    assert restarted.records()[0]['id'] == record
    assert json.loads(restarted.records()[0]['payload'])['supplier'] == 'Example Supplier'
    template = restarted.templates()[0]
    assert template['id'] == key
    source = restarted.template_path(template).read_bytes()
    assert source == out.getvalue()
    result = Document(io.BytesIO(render_document(source, {'supplier_name': 'New Supplier'}, ['supplier_name'])))
    assert result.paragraphs[0].text == 'NEW SUPPLIER'
    assert result.paragraphs[0].runs[0].bold


def test_uploaded_template_recovers_metadata_failure(tmp_path):
    server = FakeServer()
    store = SupabaseStore(server, cache_root=tmp_path)
    out = io.BytesIO()
    Document().save(out)
    server.files['retry.docx'] = out.getvalue()
    assert store.add_template('Recovered', out.getvalue(), key='retry') == 'retry'
    assert len(store.templates()) == 1
    store.add_template('Ignored seed', out.getvalue(), key='retry')
    assert store.templates()[0]['name'] == 'Recovered'


def test_login_blocks_private_store_and_disallowed_users(monkeypatch):
    calls = []
    def request(self, method, path, payload=None, **kwargs):
        calls.append(path)
        raise AssertionError('No network access expected before approved login')
    monkeypatch.setattr(SupabaseAPI, 'request', request)
    app = AppTest.from_file(str(Path(__file__).parents[1] / 'app.py'))
    app.secrets.update({'STORAGE_BACKEND': 'supabase', 'SUPABASE_URL': 'https://example.supabase.co',
                        'SUPABASE_SECRET_KEY': 'server-placeholder', 'SUPABASE_PUBLISHABLE_KEY': 'public-placeholder',
                        'ALLOWED_EMAILS': ['owner@example.com']})
    app.run()
    assert not app.exception
    assert len(app.tabs) == 0
    app.text_input[0].set_value('other@example.com')
    app.text_input[1].set_value('wrong')
    app.button[0].click().run()
    assert app.error
    assert not calls
    assert 'company_login' not in app.session_state


def test_cloud_missing_credentials_stops_without_local_fallback():
    app = AppTest.from_file(str(Path(__file__).parents[1] / 'app.py'))
    app.secrets['STORAGE_BACKEND'] = 'supabase'
    app.run()
    assert not app.exception
    assert app.error
    assert len(app.tabs) == 0


def test_company_login_expiry_and_logout(monkeypatch, tmp_path):
    import cloud_access
    server = FakeServer()
    cloud = SupabaseStore(server, cache_root=tmp_path)
    monkeypatch.setattr(cloud_access, 'SupabaseStore', lambda api: cloud)
    def authenticate(self, method, path, payload=None, **kwargs):
        assert path == '/auth/v1/token?grant_type=password'
        assert payload == {'email': 'owner@example.com', 'password': 'test-password'}
        return {'access_token': 'discarded-token', 'expires_in': 3600, 'user': {'email': 'owner@example.com'}}
    monkeypatch.setattr(SupabaseAPI, 'request', authenticate)
    app = AppTest.from_file(str(Path(__file__).parents[1] / 'app.py'), default_timeout=20)
    app.secrets.update({'STORAGE_BACKEND': 'supabase', 'SUPABASE_URL': 'https://example.supabase.co',
                        'SUPABASE_SECRET_KEY': 'server-placeholder', 'SUPABASE_PUBLISHABLE_KEY': 'public-placeholder',
                        'ALLOWED_EMAILS': ['owner@example.com']})
    app.run()
    app.text_input[0].set_value('owner@example.com')
    app.text_input[1].set_value('test-password')
    app.button[0].click().run()
    assert not app.exception
    assert len(app.tabs) == 5
    assert len(cloud.catalog('SPMC/03/2026')) == 50
    assert set(app.session_state['company_login']) == {'email', 'expires_at'}
    next(button for button in app.button if button.label == 'Sign out').click().run()
    assert not app.exception
    assert not app.tabs
    app.session_state['company_login'] = {'email': 'owner@example.com', 'expires_at': 0}
    app.run()
    assert not app.tabs
    assert 'company_login' not in app.session_state


def test_api_does_not_expose_error_body_or_credentials(monkeypatch):
    from urllib.error import HTTPError
    import cloud_store
    def reject(request, timeout):
        assert request.get_header('Apikey') == 'sb_secret_example'
        assert request.get_header('Authorization') is None
        raise HTTPError(request.full_url, 401, 'credential contents', {}, io.BytesIO(b'private details'))
    monkeypatch.setattr(cloud_store, 'urlopen', reject)
    with pytest.raises(SupabaseError) as error:
        SupabaseAPI('https://example.supabase.co', 'sb_secret_example').request('GET', '/rest/v1/usecta_entries')
    assert 'private details' not in str(error.value)
    assert 'sb_secret_example' not in str(error.value)
