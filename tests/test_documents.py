import io
import json
import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest
from docx import Document
from streamlit.testing.v1 import AppTest

import documents
from documents import Store, build_context, create_starters, render_document, template_variables


@pytest.fixture
def store(tmp_path):
    instance = Store(tmp_path)
    create_starters(instance)
    return instance


def fields():
    return {key: f"Example {key}" for key in documents.FIELDS} | {"company_name": "A&B <Trading>", "currency": "LKR"}


def test_all_starters_and_repeating_bid_items(store):
    context = build_context(fields(), [{"description": "Printer & toner", "quantity": 2, "unit_price": "100.125"}, {"description": "Paper", "quantity": 3, "unit_price": "2.10"}])
    assert context["grand_total"] == "206.56"
    assert len(store.templates()) == 3
    for template in store.templates():
        source = store.template_path(template).read_bytes()
        assert set(template["required"]) <= template_variables(source)
        output = render_document(source, context, template["required"])
        doc = Document(io.BytesIO(output))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "A&B <Trading>" in text
        assert "{{" not in text
        if template["id"] == "starter-bid":
            assert len(doc.tables[0].rows) == 3
            assert doc.tables[0].rows[1].cells[1].text == "Printer & toner"
            assert doc.tables[0].rows[2].cells[4].text == "6.30"
            assert "206.56" in text
    create_starters(store)
    assert len(store.templates()) == 3


def test_missing_fields_and_empty_bid_rejected(store):
    template = next(t for t in store.templates() if t["id"] == "starter-bid")
    with pytest.raises(ValueError, match="items"):
        render_document(store.template_path(template).read_bytes(), build_context(fields(), []), template["required"])
    template = next(t for t in store.templates() if t["id"] == "starter-authorization")
    with pytest.raises(ValueError, match="Company name"):
        render_document(store.template_path(template).read_bytes(), build_context(fields() | {"company_name": ""}, []), template["required"])


@pytest.mark.parametrize("row", [
    {"description": "Item", "quantity": -1, "unit_price": 1},
    {"description": "Item", "quantity": 0, "unit_price": 1},
    {"description": "Item", "quantity": 1, "unit_price": "NaN"},
    {"description": "", "quantity": 1, "unit_price": 1},
])
def test_invalid_items_rejected(row):
    with pytest.raises(ValueError):
        build_context(fields(), [row])


def test_custom_dotx_conversion_and_strict_fields(store):
    doc = Document()
    doc.add_paragraph("Order: {{ order_number }}")
    stream = io.BytesIO()
    doc.save(stream)
    dotx = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(stream.getvalue())) as src, zipfile.ZipFile(dotx, "w") as dest:
        for member in src.infolist():
            data = src.read(member.filename)
            if member.filename == "[Content_Types].xml":
                data = data.replace(b"wordprocessingml.document.main+xml", b"wordprocessingml.template.main+xml")
            dest.writestr(member, data)
    assert template_variables(dotx.getvalue()) == {"order_number"}
    key = store.add_template("Order", dotx.getvalue(), ["order_number"])
    template = next(t for t in store.templates() if t["id"] == key)
    rendered = render_document(store.template_path(template).read_bytes(), {"order_number": "PO-123"}, template["required"])
    assert Document(io.BytesIO(rendered)).paragraphs[0].text == "Order: PO-123"


def test_persistence_and_original_template_preserved(store):
    before = {t["id"]: store.template_path(t).read_bytes() for t in store.templates()}
    store.save_profile("Company", {"company_name": "Acme"})
    store.save_profile("Company", {"company_name": "Acme Ltd"})
    payload = {"fields": fields(), "items": []}
    key = store.save_record("Bid", payload)
    reopened = Store(store.root)
    assert reopened.profiles()["Company"]["company_name"] == "Acme Ltd"
    assert json.loads(next(r for r in reopened.records() if r["id"] == key)["payload"]) == payload
    assert before == {t["id"]: reopened.template_path(t).read_bytes() for t in reopened.templates()}


@pytest.mark.skipif(not documents.pdf_available() or not shutil.which("pdftotext"), reason="LibreOffice and pdftotext required for functional PDF validation")
def test_pdf_contains_letter_and_bid_content(store, tmp_path):
    context = build_context(fields(), [{"description": "Office printer", "quantity": 2, "unit_price": 100}])
    for template in store.templates():
        word = render_document(store.template_path(template).read_bytes(), context, template["required"])
        pdf = documents.convert_pdf(word)
        assert pdf.startswith(b"%PDF-")
        target = tmp_path / (template["id"] + ".pdf")
        target.write_bytes(pdf)
        text = subprocess.run(["pdftotext", str(target), "-"], check=True, capture_output=True, text=True).stdout
        assert "A&B <Trading>" in text
        assert "{{" not in text
        if template["id"] == "starter-bid":
            assert "Office printer" in text
            assert "200.00" in text
