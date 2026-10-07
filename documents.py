"""Local persistence, Word templates, and deterministic document generation."""
from __future__ import annotations

import io
import json
import re
import shutil
import sqlite3
import subprocess
import tempfile
import uuid
import zipfile
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from docxtpl import DocxTemplate
from jinja2 import StrictUndefined
from jinja2.sandbox import SandboxedEnvironment

DATA_DIR = Path(__file__).resolve().parent / "data"
FIELDS = {
    "company_name": "Company name", "company_address": "Company address",
    "company_email": "Company email", "company_phone": "Company phone",
    "customer_name": "Customer / recipient name", "customer_address": "Customer address",
    "document_date": "Document date", "bid_reference": "Bid reference",
    "authorized_person_name": "Authorized person name",
    "authorized_person_designation": "Authorized person designation",
    "authorized_person_id": "Authorized person ID / reference",
    "signatory_name": "Signatory name", "signatory_designation": "Signatory designation",
    "agreement_subject": "Agreement subject", "agreement_terms": "Agreement terms",
    "delivery_terms": "Delivery terms", "currency": "Currency",
}
COMPUTED = {"items", "grand_total"}
CENT = Decimal("0.01")


class Store:
    def __init__(self, root: Path = DATA_DIR):
        self.root = Path(root)
        self.root.mkdir(parents=True, exist_ok=True)
        self.templates_dir = self.root / "templates"
        self.templates_dir.mkdir(exist_ok=True)
        self.db = self.root / "records.sqlite3"
        with self.connect() as con:
            con.executescript("""
                CREATE TABLE IF NOT EXISTS profiles (
                    name TEXT PRIMARY KEY, payload TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS records (
                    id TEXT PRIMARY KEY, name TEXT NOT NULL, payload TEXT NOT NULL,
                    created_at TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS templates (
                    id TEXT PRIMARY KEY, name TEXT NOT NULL, filename TEXT NOT NULL,
                    required TEXT NOT NULL, created_at TEXT NOT NULL);
            """)

    def connect(self):
        con = sqlite3.connect(self.db)
        con.row_factory = sqlite3.Row
        return con

    def profiles(self):
        with self.connect() as con:
            return {r["name"]: json.loads(r["payload"]) for r in con.execute("SELECT * FROM profiles ORDER BY name")}

    def save_profile(self, name, payload):
        if not name.strip():
            raise ValueError("Give the profile a name.")
        with self.connect() as con:
            con.execute("INSERT INTO profiles VALUES (?, ?) ON CONFLICT(name) DO UPDATE SET payload=excluded.payload", (name.strip(), json.dumps(payload)))

    def records(self):
        with self.connect() as con:
            return [dict(r) for r in con.execute("SELECT * FROM records ORDER BY created_at DESC")]

    def save_record(self, name, payload):
        if not name.strip():
            raise ValueError("Give the project a name.")
        key = uuid.uuid4().hex
        with self.connect() as con:
            con.execute("INSERT INTO records VALUES (?, ?, ?, ?)", (key, name.strip(), json.dumps(payload), datetime.now(timezone.utc).isoformat()))
        return key

    def templates(self):
        with self.connect() as con:
            return [dict(r) | {"required": json.loads(r["required"])} for r in con.execute("SELECT * FROM templates ORDER BY name")]

    def template_path(self, template):
        return self.templates_dir / template["filename"]

    def add_template(self, name, content, required=(), key=None):
        if not name.strip():
            raise ValueError("Give the template a name.")
        validate_word(content)
        variables = template_variables(content)
        for field in variables - set(FIELDS) - COMPUTED:
            if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_]*", field):
                raise ValueError(f"Unsupported field name: {field}")
        if set(required) - variables:
            raise ValueError("Required fields must appear in the template.")
        key = key or uuid.uuid4().hex
        path = self.templates_dir / f"{key}.docx"
        path.write_bytes(content)
        with self.connect() as con:
            con.execute("INSERT INTO templates VALUES (?, ?, ?, ?, ?)", (key, name.strip(), path.name, json.dumps(list(required)), datetime.now(timezone.utc).isoformat()))
        return key


def validate_word(content):
    if len(content) > 10 * 1024 * 1024:
        raise ValueError("Templates must be smaller than 10 MB.")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError("The expanded template is too large.")
            if "word/document.xml" not in archive.namelist():
                raise ValueError("Upload a valid DOCX or DOTX Word document.")
            if any("vbaproject" in x.lower() for x in archive.namelist()):
                raise ValueError("Macro-enabled templates are not supported.")
    except zipfile.BadZipFile as exc:
        raise ValueError("Upload a valid DOCX or DOTX Word document.") from exc


def normalize_word(content):
    """Turn a DOTX package into a DOCX package without altering its layout."""
    validate_word(content)
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content)) as src, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as dest:
        for member in src.infolist():
            data = src.read(member.filename)
            if member.filename == "[Content_Types].xml":
                data = data.replace(b"application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml", b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml")
            dest.writestr(member, data)
    return output.getvalue()


def template_variables(content):
    tpl = DocxTemplate(io.BytesIO(normalize_word(content)))
    return tpl.get_undeclared_template_variables(jinja_env=SandboxedEnvironment())


def money(value):
    try:
        number = Decimal(str(value))
        if not number.is_finite() or number < 0:
            raise ValueError("Amounts and quantities must be finite and non-negative.")
        return number
    except (InvalidOperation, TypeError) as exc:
        raise ValueError("Enter valid numeric quantities and prices.") from exc


def build_context(fields, rows):
    context = {key: str(value or "") for key, value in fields.items()}
    items = []
    total = Decimal(0)
    for row in rows:
        description = str(row.get("description") or "").strip()
        if not description:
            if money(row.get("quantity", 0)) or money(row.get("unit_price", 0)):
                raise ValueError("Each priced item needs a description.")
            continue
        quantity = money(row.get("quantity", 0))
        if quantity <= 0:
            raise ValueError(f"Quantity must be greater than zero for {description}.")
        price = money(row.get("unit_price", 0)).quantize(CENT, rounding=ROUND_HALF_UP)
        subtotal = (quantity * price).quantize(CENT, rounding=ROUND_HALF_UP)
        total += subtotal
        items.append({"number": len(items) + 1, "description": description,
                      "quantity": format(quantity.normalize(), "f"), "unit_price": f"{price:,.2f}", "total": f"{subtotal:,.2f}"})
    context.update(items=items, grand_total=f"{total:,.2f}")
    return context


def render_document(content, context, required=()):
    missing = [x for x in required if not context.get(x) or (isinstance(context[x], str) and not context[x].strip())]
    if missing:
        raise ValueError("Complete required fields: " + ", ".join(FIELDS.get(x, x) for x in missing))
    template = DocxTemplate(io.BytesIO(normalize_word(content)))
    template.render(context, jinja_env=SandboxedEnvironment(undefined=StrictUndefined), autoescape=True)
    output = io.BytesIO()
    template.save(output)
    return output.getvalue()


def pdf_available():
    local = DATA_DIR / "pdf-runtime" / "bin" / "libreoffice"
    return shutil.which("libreoffice") or shutil.which("soffice") or (str(local) if local.is_file() else None)


def convert_pdf(content):
    executable = pdf_available()
    if not executable:
        raise ValueError("PDF export needs LibreOffice. Install it, or download DOCX and export to PDF in Word.")
    with tempfile.TemporaryDirectory(prefix="usecta-pdf-") as temp:
        root = Path(temp)
        source = root / "document.docx"
        source.write_bytes(content)
        result = subprocess.run([executable, f"-env:UserInstallation={(root / 'profile').as_uri()}", "--headless", "--convert-to", "pdf", "--outdir", str(root), str(source)], capture_output=True, timeout=90)
        target = root / "document.pdf"
        if result.returncode or not target.exists():
            raise ValueError("LibreOffice could not convert this document. Download the DOCX and check its layout in Word.")
        return target.read_bytes()


def filename(name):
    return re.sub(r"[^A-Za-z0-9_-]+", "_", name).strip("_")[:80] or "document"


def make_zip(files):
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in files.items():
            archive.writestr(name, content)
    return output.getvalue()


def create_starters(store):
    """Create examples once; they are starting points for the user's wording."""
    existing = {t["id"] for t in store.templates()}
    specs = [
        ("starter-authorization", "Authorization letter", "AUTHORIZATION LETTER", [
            "To: {{ customer_name }}", "{{ customer_address }}", "Reference: {{ bid_reference }}",
            "We, {{ company_name }}, authorize {{ authorized_person_name }}, {{ authorized_person_designation }} (ID/reference: {{ authorized_person_id }}), to represent our company in connection with the above reference.",
            "{{ delivery_terms }}"], ["company_name", "customer_name", "authorized_person_name", "authorized_person_designation", "bid_reference", "signatory_name"]),
        ("starter-agreement", "Agreement letter", "AGREEMENT LETTER", [
            "To: {{ customer_name }}", "{{ customer_address }}", "Subject: {{ agreement_subject }}",
            "This letter records the agreement between {{ company_name }} and {{ customer_name }} regarding {{ agreement_subject }}.",
            "Agreed terms", "{{ agreement_terms }}"], ["company_name", "customer_name", "agreement_subject", "agreement_terms", "signatory_name"]),
        ("starter-bid", "Bid covering letter", "BID COVERING LETTER", [
            "To: {{ customer_name }}", "{{ customer_address }}", "Bid reference: {{ bid_reference }}",
            "On behalf of {{ company_name }}, we submit the following items for your consideration.",
            "Delivery terms: {{ delivery_terms }}"], ["company_name", "customer_name", "bid_reference", "items", "signatory_name"]),
    ]
    for key, name, heading, paragraphs, required in specs:
        if key in existing:
            continue
        doc = Document()
        doc.sections[0].top_margin = Inches(0.75)
        normal = doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        doc.add_heading("{{ company_name }}", 0)
        doc.add_paragraph("{{ company_address }}\n{{ company_email }} | {{ company_phone }}")
        doc.add_paragraph("Date: {{ document_date }}")
        doc.add_heading(heading, 1)
        for paragraph in paragraphs:
            doc.add_paragraph(paragraph)
        if key == "starter-bid":
            table = doc.add_table(rows=1, cols=5)
            table.style = "Table Grid"
            for cell, label in zip(table.rows[0].cells, ["No.", "Description", "Quantity", "Unit price", "Total"]):
                cell.text = label
            table.add_row().cells[0].text = "{%tr for item in items %}"
            for cell, field in zip(table.add_row().cells, ["number", "description", "quantity", "unit_price", "total"]):
                cell.text = "{{ item." + field + " }}"
            table.add_row().cells[0].text = "{%tr endfor %}"
            doc.add_paragraph("Grand total: {{ currency }} {{ grand_total }}")
        doc.add_paragraph("Yours faithfully,\n\n________________________\n{{ signatory_name }}\n{{ signatory_designation }}\n{{ company_name }}")
        output = io.BytesIO()
        doc.save(output)
        store.add_template(name, output.getvalue(), required, key)
