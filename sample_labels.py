"""Editable two-copy product labels recreated from the supplied printout."""
import io
import re
import zipfile

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

LAYOUT_VERSION = 3
LABEL_GAP_POINTS = 5


def validate_labels(ifb_number, product_names):
    if not isinstance(ifb_number, str) or not ifb_number.strip():
        raise ValueError("Enter the IFB number.")
    ifb_number = ifb_number.strip()
    if len(ifb_number) > 80 or re.search(r"[\x00-\x1f\x7f]", ifb_number):
        raise ValueError("Use an IFB number of up to 80 characters on one line.")
    if not isinstance(product_names, (list, tuple)) or not product_names:
        raise ValueError("Enter at least one product name, one product per line.")
    if len(product_names) > 100:
        raise ValueError("Generate at most 100 products in one label run.")
    cleaned = []
    for name in product_names:
        if not isinstance(name, str) or not name.strip():
            raise ValueError("Every selected product needs a name.")
        if re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", name):
            raise ValueError("Product names contain unsupported control characters.")
        name = " ".join(name.split())
        if len(name) > 240:
            raise ValueError("Keep each product name within 240 characters so its label remains readable.")
        cleaned.append(name)
    return ifb_number, cleaned


def _spacer(doc, height=1, page_break=False):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(height)
    paragraph.paragraph_format.page_break_before = page_break
    paragraph.add_run().font.size = Pt(1)


def _cell_style(cell, bordered, horizontal_padding, vertical_padding=120):
    properties = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for edge, value in (("top", vertical_padding), ("bottom", vertical_padding), ("left", horizontal_padding), ("right", horizontal_padding)):
        element = OxmlElement("w:" + edge)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    properties.append(margins)
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement("w:" + edge)
        element.set(qn("w:val"), "single" if bordered else "nil")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), "000000")
        borders.append(element)
    properties.append(borders)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _reference_labels_document(ifb_number, product_names):
    ifb_number, product_names = validate_labels(ifb_number, product_names)
    doc = Document()
    doc.core_properties.title = "Sample labels"
    doc.core_properties.author = "USECTA (Pvt) Ltd"
    section = doc.sections[0]
    section.page_width, section.page_height = Pt(612), Pt(792)
    section.top_margin, section.bottom_margin = Pt(60), Pt(52)
    section.left_margin, section.right_margin = Pt(56.35), Pt(50.35)
    section.header_distance = section.footer_distance = Pt(0)
    grid = section._sectPr.find(qn("w:docGrid"))
    if grid is not None:
        section._sectPr.remove(grid)
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Calibri", Pt(16)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = 1.15
    _spacer(doc, 2.75)
    heights = [72, 96.2, 148.15]
    gaps = [72, 37.95]
    for index, name in enumerate(product_names):
        position = index % 3
        if index and position == 0:
            _spacer(doc, 2.75, page_break=True)
        table = doc.add_table(rows=1, cols=3)
        table.autofit = False
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        widths = [243, 19.3, 243]
        for column, width in zip(table.columns, widths):
            column.width = Pt(width)
        row = table.rows[0]
        # LibreOffice adds top/bottom cell margins to an at-least row height.
        row.height, row.height_rule = Pt(heights[position] - 12), WD_ROW_HEIGHT_RULE.AT_LEAST
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        for column_index, cell in enumerate(row.cells):
            cell.width = Pt(widths[column_index])
            _cell_style(cell, column_index != 1, 200 if position == 2 else 100)
            if column_index == 1:
                cell.paragraphs[0].add_run().font.size = Pt(1)
                continue
            first = cell.paragraphs[0]
            first.alignment = WD_ALIGN_PARAGRAPH.CENTER
            first.paragraph_format.space_after = Pt(10)
            first.paragraph_format.keep_with_next = True
            first.add_run("IFB NO : " + ifb_number)
            product = cell.add_paragraph(name)
            product.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if position < 2 and index < len(product_names) - 1:
            _spacer(doc, gaps[position])
    return _stable_document(doc)


def _stable_document(doc):
    output = io.BytesIO()
    doc.save(output)
    # Regenerated saved labels must not differ merely because ZIP member
    # timestamps reflect a different generation time.
    stable = io.BytesIO()
    with zipfile.ZipFile(output) as source, zipfile.ZipFile(stable, "w", zipfile.ZIP_DEFLATED) as target:
        for name in source.namelist():
            entry = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            entry.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(entry, source.read(name))
    return stable.getvalue()


def labels_document(ifb_number, product_names, layout_version=LAYOUT_VERSION):
    if layout_version == 1:
        return _reference_labels_document(ifb_number, product_names)
    if layout_version not in (2, LAYOUT_VERSION):
        raise ValueError("This saved label layout is not supported by this app version.")
    ifb_number, product_names = validate_labels(ifb_number, product_names)
    doc = Document()
    doc.core_properties.title = "Sample labels"
    doc.core_properties.author = "USECTA (Pvt) Ltd"
    section = doc.sections[0]
    section.page_width, section.page_height = Pt(612), Pt(792)
    section.top_margin = section.bottom_margin = Pt(24)
    section.left_margin = section.right_margin = Pt(24)
    section.header_distance = section.footer_distance = Pt(0)
    grid = section._sectPr.find(qn("w:docGrid"))
    if grid is not None:
        section._sectPr.remove(grid)
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Calibri", Pt(14)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = 1.0
    # Layout 3 uses separate one-row tables with small blank paragraphs between
    # them. A continuous table can draw unwanted border rules at page breaks.
    table = doc.add_table(rows=0, cols=3)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    widths = [276, 12, 276]
    for column, width in zip(table.columns, widths):
        column.width = Pt(width)
    for index, name in enumerate(product_names):
        if index and layout_version == 3:
            _spacer(doc, LABEL_GAP_POINTS)
            doc.paragraphs[-1].paragraph_format.keep_with_next = True
            table = doc.add_table(rows=0, cols=3)
            table.autofit = False
            table.alignment = WD_TABLE_ALIGNMENT.LEFT
            for column, width in zip(table.columns, widths):
                column.width = Pt(width)
        row = table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for column_index, cell in enumerate(row.cells):
            cell.width = Pt(widths[column_index])
            _cell_style(cell, column_index != 1, 120, vertical_padding=60)
            first = cell.paragraphs[0]
            if column_index == 1:
                first.add_run().font.size = Pt(1)
                continue
            first.alignment = WD_ALIGN_PARAGRAPH.CENTER
            first.paragraph_format.space_after = Pt(2)
            first.paragraph_format.keep_with_next = True
            first.paragraph_format.keep_together = True
            first.add_run("IFB NO : " + ifb_number).font.size = Pt(12)
            product = cell.add_paragraph(name)
            product.alignment = WD_ALIGN_PARAGRAPH.CENTER
            product.paragraph_format.keep_together = True
    # The mandatory paragraph after the table must not create a blank last page.
    _spacer(doc, 1)
    return _stable_document(doc)


def label_files(ifb_number, product_names, layout_version=LAYOUT_VERSION):
    from documents import filename
    return {filename(ifb_number)[:65] + "_Sample_labels.docx": labels_document(ifb_number, product_names, layout_version)}
