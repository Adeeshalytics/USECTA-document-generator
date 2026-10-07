"""Tender catalogues and quotations have independent data and calculations."""
from __future__ import annotations

import calendar
import csv
import io
import json
import hashlib
from pathlib import Path
from datetime import date, time
from decimal import ROUND_HALF_UP

from documents import (CENT, Store, build_context, convert_pdfs, filename, money,
                       render_document, template_variables)

SUPPLIER_FIELDS = {
    "supplier_name": "Supplier name", "supplier_country": "Country",
    "supplier_registration_number": "Registration number", "supplier_address": "Registered address",
    "supplier_signatory_name": "Supplier signatory name",
    "supplier_signatory_designation": "Supplier signatory designation",
}
TENDER_COMPUTED = {"bid_reference", "closing_date_long", "closing_date_short", "closing_time",
                   "tender_title", "bid_valid_until", "tender_items", "item_number", "item_description",
                   "agreement_item_names", "item_quantity_kg"}
QUOTE_COMPUTED = {"items", "quotation_items", "subtotal", "grand_total"}
CATALOG_COLUMNS = ["item_number", "description", "document_fee"]


def iso_date(value, label):
    try:
        return date.fromisoformat(str(value))
    except (ValueError, TypeError) as exc:
        raise ValueError(f"{label} must be a valid date (YYYY-MM-DD).") from exc


def long_date(value):
    value = iso_date(value, "Date")
    return f"{value.day} {calendar.month_name[value.month]} {value.year}"


def validate_tender(reference, payload):
    if not reference:
        raise ValueError("Enter a tender reference.")
    iso_date(payload.get("closing_date"), "Closing date")
    try:
        time.fromisoformat(str(payload.get("closing_time", "")))
    except ValueError as exc:
        raise ValueError("Enter the closing time in HH:MM format.") from exc
    if payload.get("bid_valid_until"):
        if iso_date(payload["bid_valid_until"], "Bid validity date") < iso_date(payload["closing_date"], "Closing date"):
            raise ValueError("Bid validity must extend to or beyond the closing date.")


def item_number(value):
    if value is None:
        return ""
    text = str(value).strip()
    # Excel commonly exports integer identifiers as floating point numbers.
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return text


def validate_catalog(rows):
    result, seen = [], set()
    for row in rows:
        number = item_number(row.get("item_number"))
        description = str(row.get("description") or "").strip()
        fee = str(row.get("document_fee") if row.get("document_fee") is not None else "").strip()
        if not number and not description and not fee:
            continue
        if not number or not description or not fee:
            raise ValueError("Every catalogue row needs an item number, description, and document fee (use 0 if free).")
        if number in seen:
            raise ValueError(f"Duplicate item number: {number}")
        amount = money(fee.replace(",", "")).quantize(CENT, rounding=ROUND_HALF_UP)
        quantity = str(row.get("quantity_kg") if row.get("quantity_kg") is not None else "").strip()
        if quantity:
            quantity = format(money(quantity.replace(",", "")).normalize(), "f")
        result.append({"item_number": number, "description": description, "document_fee": str(amount), "quantity_kg": quantity})
        seen.add(number)
    if not result:
        raise ValueError("Add at least one catalogue item.")
    return result


def import_catalog(content, extension):
    if len(content) > 10 * 1024 * 1024:
        raise ValueError("Catalogue files must be smaller than 10 MB.")
    if extension.lower() == ".csv":
        try:
            reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
            headers = reader.fieldnames or []
            rows = list(reader)
        except UnicodeError as exc:
            raise ValueError("Save your CSV as UTF-8, or upload an XLSX file.") from exc
    elif extension.lower() == ".xlsx":
        from openpyxl import load_workbook
        # Bound the expanded archive before handing an XLSX package to openpyxl.
        import zipfile
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError("The expanded workbook is too large.")
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            values = book.active.iter_rows(values_only=True)
            headers = [str(x or "").strip() for x in next(values, ())]
            rows = [dict(zip(headers, row)) for row in values]
        finally:
            book.close()
    else:
        raise ValueError("Upload a CSV or XLSX catalogue.")
    if not set(CATALOG_COLUMNS) <= set(headers):
        raise ValueError("Use these column headers: item_number, description, document_fee.")
    return validate_catalog(rows)


def choose_items(store, reference, numbers):
    catalog = {row["item_number"]: row for row in store.catalog(reference)}
    selected = []
    seen = set()
    for number in numbers:
        number = item_number(number)
        if number in seen:
            raise ValueError(f"Item {number} was selected twice.")
        if number not in catalog:
            raise ValueError(f"Item {number} is not in tender {reference}.")
        selected.append(catalog[number])
        seen.add(number)
    return selected


def tender_context(reference, tender, rows, supplier, document_date, effective_date):
    validate_tender(reference, tender)
    closing = iso_date(tender["closing_date"], "Closing date")
    clock = time.fromisoformat(tender["closing_time"])
    description_list = [row["description"] for row in rows]
    names = (", ".join(description_list[:-1]) + " and " + description_list[-1]
             if len(description_list) > 1 else "".join(description_list))
    context = {
        "bid_reference": reference, "document_date": long_date(document_date),
        "agreement_effective_date": long_date(effective_date),
        "closing_date_long": long_date(closing), "closing_date_short": closing.strftime("%d.%m.%Y"),
        "closing_time": f"{clock.hour % 12 or 12}.{clock.minute:02d} {'a.m.' if clock.hour < 12 else 'p.m.'}",
        "bid_valid_until": iso_date(tender["bid_valid_until"], "Bid validity date").strftime("%d.%m.%Y") if tender.get("bid_valid_until") else "",
        "tender_title": tender.get("tender_title", ""),
        "agreement_item_names": names, "item_number": "", "item_description": "", "item_quantity_kg": "",
        "tender_items": [{"item_number": row["item_number"], "description": row["description"],
                          "document_fee": f"{money(row['document_fee']):,.2f}", "quantity_kg": row.get("quantity_kg", "")} for row in rows],
    }
    context.update({key: supplier.get(key, "") for key in SUPPLIER_FIELDS})
    return context


def tender_jobs(store, template_ids, context):
    templates = {t["id"]: t for t in store.templates()}
    jobs = []
    for key in template_ids:
        template = templates[key]
        if template["workflow"] != "tender":
            raise ValueError("Select only tender templates in the tender workflow.")
        variables = template_variables(store.template_path(template).read_bytes())
        if {"item_number", "item_description"} & variables and template["scope"] != "item":
            raise ValueError(f"{template['name']} contains a single-item field. Set its scope to 'One document per item' in Templates.")
        contexts = [context]
        if template["scope"] == "item":
            if not context["tender_items"]:
                raise ValueError("Select at least one item for per-item documents.")
            contexts = [context | {"item_number": row["item_number"], "item_description": row["description"], "item_quantity_kg": row["quantity_kg"]} for row in context["tender_items"]]
        for current in contexts:
            suffix = f"_item_{filename(current['item_number'])}" if template["scope"] == "item" else ""
            stem = f"{filename(context['bid_reference'])}_{filename(template['name'])}{suffix}_{key[:8]}"
            jobs.append({"template_id": key, "context": current, "filename": stem})
    return jobs


def quotation_context(fields, rows):
    # Tender fields are deliberately excluded from quotation data.
    clean = {k: v for k, v in fields.items() if k not in TENDER_COMPUTED and k not in SUPPLIER_FIELDS}
    context = build_context(clean, rows)
    nonblank = [r for r in rows if str(r.get("description") or "").strip()]
    for generated, source in zip(context["items"], nonblank):
        generated["unit"] = str(source.get("unit") or "").strip()
    context["quotation_items"] = context["items"]
    context["subtotal"] = context["grand_total"]
    return context


def quotation_jobs(store, template_ids, context):
    templates = {t["id"]: t for t in store.templates()}
    jobs = []
    for key in template_ids:
        template = templates[key]
        if template["workflow"] != "quotation":
            raise ValueError("Select only quotation templates in the quotation workflow.")
        variables = template_variables(store.template_path(template).read_bytes())
        if variables & TENDER_COMPUTED:
            raise ValueError("Quotation templates cannot use tender fields or tender item lists.")
        jobs.append({"template_id": key, "context": context,
                     "filename": filename(context.get("quotation_number", "Quotation")) + "_" + filename(template["name"]) + "_" + key[:8]})
    return jobs


def generate_files(store, jobs, with_pdf=False):
    if not jobs:
        raise ValueError("Select at least one document template.")
    templates = {t["id"]: t for t in store.templates()}
    files = {}
    # Render every Word file before attempting PDF conversion. A failure produces no partial download set.
    for job in jobs:
        template = templates.get(job["template_id"])
        if template is None:
            raise ValueError("A saved template version is missing. Restore it from your data backup.")
        try:
            content = render_document(store.template_path(template).read_bytes(), job["context"], template["required"])
        except Exception as exc:
            raise ValueError(f"{template['name']}: {exc}") from exc
        path = job["filename"] + ".docx"
        if path in files:
            raise ValueError("Two output filenames are identical. Change the template names or item identifiers.")
        files[path] = content
    if with_pdf:
        files.update(convert_pdfs(files))
    return files


def create_quotation_template(store: Store):
    if any(t["workflow"] == "quotation" or t["id"] == "usecta-quotation-starter" for t in store.templates()):
        return
    from docx import Document
    doc = Document()
    doc.add_heading("USECTA (Pvt) Ltd", 0)
    doc.add_paragraph("33/5 Boruppa Road, Pallegunnepana, Gunnepana Kandy")
    doc.add_heading("QUOTATION {{ quotation_number }}", 1)
    doc.add_paragraph("Date: {{ document_date }}\nValid until: {{ quotation_expiry_date }}")
    doc.add_paragraph("To: {{ customer_name }}\n{{ customer_address }}")
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, ["No.", "Description", "Qty", "Unit", "Rate", "Amount"]):
        cell.text = label
    table.add_row().cells[0].text = "{%tr for item in quotation_items %}"
    for cell, field in zip(table.add_row().cells, ["number", "description", "quantity", "unit", "unit_price", "total"]):
        cell.text = "{{ item." + field + " }}"
    table.add_row().cells[0].text = "{%tr endfor %}"
    doc.add_paragraph("Total: {{ currency }} {{ grand_total }}")
    doc.add_paragraph("Shipment mode: {{ shipment_mode }}\nContainer: {{ container_details }}\nQuantity: {{ shipment_quantity }}\nPackaging: {{ packaging_details }}\nPayment terms: {{ payment_terms }}\n{{ quotation_notes }}")
    output = io.BytesIO()
    doc.save(output)
    required = ["quotation_number", "document_date", "quotation_expiry_date", "customer_name", "quotation_items", "currency"]
    store.add_template("USECTA quotation (starter)", output.getvalue(), required, key="usecta-quotation-starter", workflow="quotation")


def prepare_workspace(store, project_root=None):
    """Install supplied assets once; never replace user-edited catalogues or templates."""
    root = Path(project_root) if project_root else Path(__file__).resolve().parent
    metadata = root / "catalogues/SPMC_03_2026.json"
    catalogue = root / "catalogues/SPMC_03_2026.csv"
    if metadata.exists() and catalogue.exists():
        details = json.loads(metadata.read_text())
        reference = details.pop("reference")
        if reference not in store.tenders():
            rows = import_catalog(catalogue.read_bytes(), ".csv")
            store.save_tender(reference, details)
            store.save_catalog(reference, rows)
    folder = root / "company_templates"
    manifest = folder / "manifest.json"
    if manifest.exists():
        known = {t["id"] for t in store.templates()}
        for entry in json.loads(manifest.read_text()):
            path = folder / entry["filename"]
            if path.resolve().parent != folder.resolve() or path.suffix.lower() != ".docx":
                raise ValueError("Company template manifest contains an invalid file path.")
            content = path.read_bytes()
            key = "usecta-" + hashlib.sha256(content).hexdigest()[:20]
            if key not in known:
                variables = sorted(template_variables(content))
                store.add_template(entry["name"], content, variables, key=key, workflow="tender", scope=entry["scope"])
                known.add(key)
    create_quotation_template(store)
